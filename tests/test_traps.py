"""Trap simulator — proves the verifier defeats the held-out adversarial classes.

The graded mailbox adds attacks a naive "file arrived → Complete" agent fails: wrong period,
wrong entity, missing-invoice false completion, hidden Excel tabs, and misleading filenames.
Each test synthesizes that trap and asserts the agent downgrades correctly, with a reason.
Self-contained (builds its own documents), so it runs with no bundle or key.
"""

from __future__ import annotations

import io
from datetime import date

from pbc_agent.config.engagement import Engagement, Entity
from pbc_agent.model.assessment import Status
from pbc_agent.model.criteria import (
    CountCriterion,
    EntityCriterion,
    PBCItem,
    PeriodCriterion,
    PeriodKind,
    Priority,
    SignatureCriterion,
)
from pbc_agent.model.documents import Document, SourceType
from pbc_agent.tools_impl.extract import extract_fields
from pbc_agent.tools_impl.parse import parse_document
from pbc_agent.tools_impl.verify import verify_item

ENG = Engagement(
    client_name="Northwind Beverages, Inc.", fiscal_year_end=date(2026, 6, 30),
    entities=[Entity("Northwind Beverages, Inc.", "parent"),
              Entity("Northwind Distribution LLC", "sub"),
              Entity("Cascade Cold Brew LLC", "sub")])


def _pdf_doc(text: str, filename="doc.pdf") -> Document:
    d = Document(doc_id=filename, filename=filename, container_chain=[filename],
                 sniffed_type=SourceType.PDF_NATIVE)
    d.text = text
    d.pages = [text]
    d.parsed = True
    extract_fields(d, ENG)
    return d


def _item(*criteria) -> PBCItem:
    return PBCItem(id="PBC-X", category="Test", priority=Priority.HIGH,
                   description="trap item", criteria=list(criteria))


def test_trap_wrong_period():
    # A prior-year document offered for a current-year request.
    doc = _pdf_doc("Consolidated balances as of June 30, 2025.", "TB_FY26.pdf")
    item = _item(PeriodCriterion(raw="as_of=2026-06-30", kind=PeriodKind.AS_OF,
                                 end=date(2026, 6, 30), start=date(2026, 6, 30)))
    a = verify_item(item, [doc], ENG)
    assert a.status is Status.INSUFFICIENT
    assert "2025" in a.reasoning and "2026" in a.reasoning   # names the mismatch


def test_trap_wrong_entity():
    # A "consolidated" TB that only shows one subsidiary — parent + others missing.
    doc = _pdf_doc("Trial balance for Cascade Cold Brew LLC as of 2026-06-30.")
    item = _item(EntityCriterion(raw="entity=consolidated", scope_word="consolidated",
                                 required=ENG.entity_names))
    a = verify_item(item, [doc], ENG)
    assert a.status is Status.INSUFFICIENT
    assert any("missing" in oi.lower() for oi in a.open_items)


def test_trap_missing_invoice_false_completion():
    # Bundle claims a sample of 30 but only 22 documents are actually present.
    docs = [_pdf_doc(f"Invoice {i} dated 2026-06-2{i%10}", f"inv_{i}.pdf") for i in range(22)]
    docs[0].text += "\nAP cutoff sample of 30 invoices enclosed."
    item = _item(CountCriterion(raw="sample_size=30", sample_size=30, unit="invoices"))
    a = verify_item(item, docs, ENG)
    assert a.status is Status.UNDER_REVIEW           # partial, not a rubber-stamped Complete
    assert any("missing" in oi.lower() for oi in a.open_items)


def test_trap_misleading_filename_unsigned():
    # File named "..._signed.pdf" but the content is a DRAFT placeholder.
    doc = _pdf_doc("Management representation letter (DRAFT)\n/s/ [TO BE SIGNED]",
                   "MgmtRep_signed_FINAL.pdf")
    item = _item(SignatureCriterion(raw="signed=True", signed=True))
    a = verify_item(item, [doc], ENG)
    assert a.status is Status.UNDER_REVIEW            # not Complete despite the filename
    assert any("draft" in oi.lower() for oi in a.open_items)


def test_trap_hidden_excel_tab_is_read():
    # A hidden worksheet must still be parsed (data can't be concealed from the verifier).
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.active["A1"] = "Nothing to see"
    hidden = wb.create_sheet("Adjustments")
    hidden["A1"] = "Secret AJE 2026-06-30 for Cascade Cold Brew LLC"
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)

    d = Document(doc_id="x", filename="TB.xlsx", container_chain=["TB.xlsx"],
                 sniffed_type=SourceType.XLSX, raw_bytes=buf.getvalue())
    parse_document(d)
    states = {s.name: s.state for s in d.sheets}
    assert states.get("Adjustments") == "hidden"      # detected as hidden...
    assert "Secret AJE" in d.text                       # ...but still read


if __name__ == "__main__":
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"ok  {fn.__name__}")
    print(f"\n{len(fns)} passed")
